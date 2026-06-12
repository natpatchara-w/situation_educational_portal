import json
import uuid
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from django.contrib.auth.models import User
from django.contrib.auth.models import Permission
from django.core.exceptions import ImproperlyConfigured, ValidationError
from django.core.files.base import ContentFile
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase, override_settings
from django.utils import timezone
from django.urls import reverse
from docx import Document

from . import education_chat
from .checklist_generator import extract_docx_text, generate_checklist_payload, normalize_checklist_payload, redact_sensitive_text
from .cleanup import delete_expired_checklist_jobs
from .models import ChatSource, ChecklistJob, OpenAISettings, Resource, ThrottleRecord
from .pdf_utils import build_simple_pdf
from .throttling import client_ip
from volunteer_portal.database import build_database_config, database_config_from_url


def build_simple_docx():
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:body><w:p><w:r><w:t>Volunteer checklist</w:t></w:r></w:p></w:body>"
        "</w:document>"
    )
    content_types_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        "</Types>"
    )

    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types_xml)
        archive.writestr("word/document.xml", document_xml)
    return buffer.getvalue()


def build_readable_docx(text="Volunteer checklist", include_heading=True):
    document = Document()
    if include_heading:
        document.add_heading("Event Concept Note", level=1)
    if text:
        document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def checklist_job_payload(upload, language="en"):
    return {"concept_note": upload, "ai_processing_acknowledged": "true", "language": language}


class FakeResponseHeaders:
    def __init__(self, content_type):
        self.content_type = content_type

    def get(self, name, default=None):
        if name.lower() == "content-type":
            return self.content_type
        return default

    def get_content_charset(self):
        return "utf-8"


class FakeWebsiteResponse:
    def __init__(self, body, content_type="text/html; charset=utf-8"):
        self.body = body.encode("utf-8") if isinstance(body, str) else body
        self.headers = FakeResponseHeaders(content_type)

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, traceback):
        return False

    def read(self, size=-1):
        return self.body if size < 0 else self.body[:size]


class DatabaseConfigTests(TestCase):
    def test_local_database_defaults_to_sqlite(self):
        config = build_database_config({}, Path("/tmp/app/backend"), is_production=False)

        self.assertEqual(config["ENGINE"], "django.db.backends.sqlite3")
        self.assertEqual(config["NAME"], Path("/tmp/app/backend") / "db.sqlite3")

    def test_production_requires_database_url(self):
        with self.assertRaises(ImproperlyConfigured):
            build_database_config({}, Path("/tmp/app/backend"), is_production=True)

    def test_postgres_database_url_is_supported(self):
        config = build_database_config(
            {
                "DJANGO_DATABASE_URL": "postgres://portal:secret@db.example.org:5432/volunteers?sslmode=verify-full",
            },
            Path("/tmp/app/backend"),
            is_production=True,
        )

        self.assertEqual(config["ENGINE"], "django.db.backends.postgresql")
        self.assertEqual(config["NAME"], "volunteers")
        self.assertEqual(config["USER"], "portal")
        self.assertEqual(config["PASSWORD"], "secret")
        self.assertEqual(config["HOST"], "db.example.org")
        self.assertEqual(config["PORT"], "5432")
        self.assertEqual(config["OPTIONS"]["sslmode"], "verify-full")

    def test_sqlite_url_can_be_used_locally(self):
        config = database_config_from_url("sqlite:///custom.sqlite3", Path("/tmp/app/backend"))

        self.assertEqual(config["ENGINE"], "django.db.backends.sqlite3")
        self.assertEqual(config["NAME"], Path("/tmp/app/backend/custom.sqlite3"))


class ResourceApiTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="volunteer", password="test-password")
        self.checklist = Resource.objects.create(
            title="Beach Event Checklist",
            description="Packing and safety list",
            category=Resource.Category.CHECKLIST,
            pdf_file=SimpleUploadedFile("checklist.pdf", build_simple_pdf("Checklist", "Test PDF"), content_type="application/pdf"),
        )
        self.educational = Resource.objects.create(
            title="Tsunami Education Guide",
            description="Learning material for students",
            category=Resource.Category.EDUCATIONAL,
            pdf_file=SimpleUploadedFile("guide.pdf", build_simple_pdf("Guide", "Test PDF"), content_type="application/pdf"),
        )
        self.generator_only = Resource.objects.create(
            title="Generator Handbook",
            category=Resource.Category.EDUCATIONAL,
            access_level=Resource.AccessLevel.CHECKLIST_GENERATORS,
            pdf_file=SimpleUploadedFile("generator.pdf", build_simple_pdf("Generator", "Test PDF"), content_type="application/pdf"),
        )
        self.staff_only = Resource.objects.create(
            title="Staff Operations Note",
            category=Resource.Category.EDUCATIONAL,
            access_level=Resource.AccessLevel.STAFF,
            pdf_file=SimpleUploadedFile("staff.pdf", build_simple_pdf("Staff", "Test PDF"), content_type="application/pdf"),
        )
        Resource.objects.create(
            title="Hidden Draft",
            category=Resource.Category.CHECKLIST,
            is_active=False,
            pdf_file=SimpleUploadedFile("draft.pdf", build_simple_pdf("Draft", "Test PDF"), content_type="application/pdf"),
        )

    def test_anonymous_users_cannot_list_resources(self):
        response = self.client.get(reverse("api-resource-list"))

        self.assertEqual(response.status_code, 401)

    def test_logged_in_users_can_list_active_resources(self):
        self.client.login(username="volunteer", password="test-password")

        response = self.client.get(reverse("api-resource-list"))

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(len(payload["resources"]), 2)
        uuid.UUID(payload["resources"][0]["id"])

    def test_generator_permission_can_list_generator_resources(self):
        self.user.user_permissions.add(Permission.objects.get(codename="can_generate_checklist"))
        self.client.login(username="volunteer", password="test-password")

        response = self.client.get(reverse("api-resource-list"))

        self.assertEqual(response.status_code, 200)
        titles = {resource["title"] for resource in response.json()["resources"]}
        self.assertIn("Generator Handbook", titles)
        self.assertNotIn("Staff Operations Note", titles)

    def test_staff_can_list_staff_resources(self):
        self.user.is_staff = True
        self.user.save()
        self.client.login(username="volunteer", password="test-password")

        response = self.client.get(reverse("api-resource-list"))

        self.assertEqual(response.status_code, 200)
        titles = {resource["title"] for resource in response.json()["resources"]}
        self.assertIn("Generator Handbook", titles)
        self.assertIn("Staff Operations Note", titles)

    def test_cors_allows_configured_frontend_origin(self):
        response = self.client.options(reverse("api-resource-list"), HTTP_ORIGIN="http://localhost:5173")

        self.assertEqual(response["Access-Control-Allow-Origin"], "http://localhost:5173")
        self.assertEqual(response["Access-Control-Allow-Credentials"], "true")

    def test_cors_blocks_unconfigured_origin(self):
        response = self.client.options(reverse("api-resource-list"), HTTP_ORIGIN="https://evil.example")

        self.assertNotIn("Access-Control-Allow-Origin", response)

    @override_settings(CSP_ENABLED=True, CSP_REPORT_ONLY=False, CSP_POLICY="default-src 'self'")
    def test_csp_header_can_be_enabled(self):
        self.client.login(username="volunteer", password="test-password")

        response = self.client.get(reverse("api-resource-list"))

        self.assertEqual(response["Content-Security-Policy"], "default-src 'self'")

    @override_settings(CSP_ENABLED=True, CSP_REPORT_ONLY=True, CSP_POLICY="default-src 'self'")
    def test_csp_report_only_header_can_be_enabled(self):
        self.client.login(username="volunteer", password="test-password")

        response = self.client.get(reverse("api-resource-list"))

        self.assertEqual(response["Content-Security-Policy-Report-Only"], "default-src 'self'")

    def test_search_and_category_filters_resources(self):
        self.client.login(username="volunteer", password="test-password")

        response = self.client.get(reverse("api-resource-list"), {"search": "tsunami", "category": "educational"})

        self.assertEqual(response.status_code, 200)
        resources = response.json()["resources"]
        self.assertEqual(len(resources), 1)
        self.assertEqual(resources[0]["title"], "Tsunami Education Guide")

    @override_settings(RESOURCE_SEARCH_MAX_LENGTH=6)
    def test_search_rejects_excessive_length(self):
        self.client.login(username="volunteer", password="test-password")

        response = self.client.get(reverse("api-resource-list"), {"search": "tsunami"})

        self.assertEqual(response.status_code, 400)

    def test_download_serves_active_pdf(self):
        self.client.login(username="volunteer", password="test-password")

        response = self.client.get(reverse("api-resource-download", args=[self.checklist.public_id]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn("attachment", response["Content-Disposition"])

    def test_resource_download_emits_security_audit_log(self):
        self.client.login(username="volunteer", password="test-password")

        with self.assertLogs("resources.security", level="INFO") as logs:
            response = self.client.get(reverse("api-resource-download", args=[self.checklist.public_id]))

        self.assertEqual(response.status_code, 200)
        self.assertIn("event=resource_download", logs.output[0])

    def test_download_serves_active_docx(self):
        docx = Resource.objects.create(
            title="Editable Volunteer Checklist",
            description="Word document version",
            category=Resource.Category.CHECKLIST,
            pdf_file=SimpleUploadedFile(
                "editable-checklist.docx",
                build_simple_docx(),
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        )
        self.client.login(username="volunteer", password="test-password")

        response = self.client.get(reverse("api-resource-download", args=[docx.public_id]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertIn("editable-checklist", response["Content-Disposition"])
        self.assertIn(".docx", response["Content-Disposition"])

    def test_download_rejects_missing_resource(self):
        self.client.login(username="volunteer", password="test-password")

        response = self.client.get(reverse("api-resource-download", args=[uuid.uuid4()]))

        self.assertEqual(response.status_code, 404)

    def test_download_rejects_resource_outside_user_access_level(self):
        self.client.login(username="volunteer", password="test-password")

        response = self.client.get(reverse("api-resource-download", args=[self.staff_only.public_id]))

        self.assertEqual(response.status_code, 404)

    def test_pdf_upload_must_have_pdf_signature(self):
        resource = Resource(
            title="Bad PDF",
            category=Resource.Category.CHECKLIST,
            pdf_file=SimpleUploadedFile("bad.pdf", b"not a pdf", content_type="application/pdf"),
        )

        with self.assertRaises(ValidationError):
            resource.full_clean()

    def test_docx_upload_must_be_valid_docx_zip(self):
        resource = Resource(
            title="Bad DOCX",
            category=Resource.Category.CHECKLIST,
            pdf_file=SimpleUploadedFile("bad.docx", b"not a docx", content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        )

        with self.assertRaises(ValidationError):
            resource.full_clean()

    @override_settings(RESOURCE_MAX_UPLOAD_BYTES=8)
    def test_resource_upload_rejects_oversized_file(self):
        resource = Resource(
            title="Large PDF",
            category=Resource.Category.CHECKLIST,
            pdf_file=SimpleUploadedFile("large.pdf", b"%PDF-" + b"x" * 20, content_type="application/pdf"),
        )

        with self.assertRaises(ValidationError):
            resource.full_clean()

    @override_settings(UPLOAD_SCANNING_REQUIRED=True, UPLOAD_SCAN_COMMAND="")
    def test_resource_upload_requires_scanner_when_configured(self):
        resource = Resource(
            title="Needs Scan",
            category=Resource.Category.CHECKLIST,
            pdf_file=SimpleUploadedFile("scan.pdf", build_simple_pdf("Scan", "Test PDF"), content_type="application/pdf"),
        )

        with self.assertRaises(ValidationError):
            resource.full_clean()


class ChecklistGeneratorApiTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="volunteer", password="test-password")

    def test_anonymous_users_cannot_generate_checklist(self):
        response = self.client.post(reverse("api-checklist-generate"))

        self.assertEqual(response.status_code, 401)

    def test_logged_in_user_without_permission_cannot_create_checklist_job(self):
        self.client.login(username="volunteer", password="test-password")
        upload = SimpleUploadedFile(
            "concept-note.docx",
            build_readable_docx("Marimba workshop for village students."),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        response = self.client.post(reverse("api-checklist-job-create"), checklist_job_payload(upload))

        self.assertEqual(response.status_code, 403)

    @override_settings(AI_CHECKLIST_GENERATION_ENABLED=False)
    def test_checklist_job_create_respects_ai_disable_flag(self):
        self.client.login(username="volunteer", password="test-password")
        self.user.user_permissions.add(Permission.objects.get(codename="can_generate_checklist"))
        upload = SimpleUploadedFile(
            "concept-note.docx",
            build_readable_docx("Marimba workshop for village students."),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        response = self.client.post(reverse("api-checklist-job-create"), checklist_job_payload(upload))

        self.assertEqual(response.status_code, 503)

    def test_checklist_job_create_requires_ai_acknowledgement(self):
        self.client.login(username="volunteer", password="test-password")
        self.user.user_permissions.add(Permission.objects.get(codename="can_generate_checklist"))
        upload = SimpleUploadedFile(
            "concept-note.docx",
            build_readable_docx("Marimba workshop for village students."),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        response = self.client.post(reverse("api-checklist-job-create"), {"concept_note": upload})

        self.assertEqual(response.status_code, 400)

    def test_checklist_job_create_rejects_unsupported_language(self):
        self.client.login(username="volunteer", password="test-password")
        self.user.user_permissions.add(Permission.objects.get(codename="can_generate_checklist"))
        upload = SimpleUploadedFile(
            "concept-note.docx",
            build_readable_docx("Marimba workshop for village students."),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        response = self.client.post(reverse("api-checklist-job-create"), checklist_job_payload(upload, language="fr"))

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["detail"], "Unsupported language.")
        self.assertFalse(ChecklistJob.objects.exists())

    def test_checklist_job_create_stores_and_serializes_language(self):
        self.client.login(username="volunteer", password="test-password")
        self.user.user_permissions.add(Permission.objects.get(codename="can_generate_checklist"))
        upload = SimpleUploadedFile(
            "concept-note.docx",
            build_readable_docx("Marimba workshop for village students."),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        with patch("resources.tasks.generate_checklist_job.delay"):
            response = self.client.post(reverse("api-checklist-job-create"), checklist_job_payload(upload, language="id"))

        self.assertEqual(response.status_code, 202)
        self.assertEqual(response.json()["job"]["language"], "id")
        self.assertEqual(ChecklistJob.objects.get().language, ChecklistJob.Language.INDONESIAN)

    @override_settings(CHECKLIST_JOB_THROTTLE_LIMIT=1)
    def test_checklist_job_create_is_throttled(self):
        self.client.login(username="volunteer", password="test-password")
        self.user.user_permissions.add(Permission.objects.get(codename="can_generate_checklist"))
        OpenAISettings.objects.create(api_key="sk-test")

        for _ in range(2):
            upload = SimpleUploadedFile(
                "concept-note.docx",
                build_readable_docx("Marimba workshop for village students."),
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
            with patch("resources.tasks.generate_checklist_job.delay"):
                response = self.client.post(reverse("api-checklist-job-create"), checklist_job_payload(upload))

        self.assertEqual(response.status_code, 429)

    def test_extract_docx_text_reads_paragraphs(self):
        upload = SimpleUploadedFile(
            "concept-note.docx",
            build_readable_docx("Marimba workshop for village students."),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        text = extract_docx_text(upload)

        self.assertIn("Event Concept Note", text)
        self.assertIn("Marimba workshop", text)

    def test_redact_sensitive_text_masks_obvious_personal_data_and_tokens(self):
        redacted = redact_sensitive_text("Email jane@example.org, call +1 555-111-2222, api_key=sk_secretvalue12345")

        self.assertIn("[redacted-email]", redacted)
        self.assertIn("[redacted-phone]", redacted)
        self.assertIn("[redacted-secret]", redacted)

    @override_settings(CHECKLIST_MAX_SECTIONS=1, CHECKLIST_MAX_ITEMS_PER_SECTION=1)
    def test_normalize_checklist_payload_bounds_model_output(self):
        payload = normalize_checklist_payload(
            {
                "event_title": "x" * 300,
                "source_note": "y" * 700,
                "sections": [
                    {"title": "First", "items": ["a", "b"]},
                    {"title": "Second", "items": ["c"]},
                ],
            }
        )

        self.assertEqual(len(payload["event_title"]), 160)
        self.assertEqual(len(payload["source_note"]), 500)
        self.assertEqual(len(payload["sections"]), 1)
        self.assertEqual(payload["sections"][0]["items"], ["a"])

    def test_generate_checklist_redacts_prompt_before_openai_call(self):
        output_text = json.dumps(
            {
                "event_title": "Safe Workshop",
                "source_note": "Generated.",
                "sections": [{"title": "Quick checklist", "items": ["Bring water."]}],
            }
        )

        with patch("resources.checklist_generator.OpenAI") as openai:
            openai.return_value.responses.create.return_value = SimpleNamespace(output_text=output_text)
            generate_checklist_payload("Contact jane@example.org and api_key=sk_secretvalue12345", "sk-test")

        openai.assert_called_once()
        self.assertEqual(openai.call_args.kwargs["timeout"], 180)
        self.assertEqual(openai.call_args.kwargs["max_retries"], 0)
        self.assertEqual(openai.return_value.responses.create.call_args.kwargs["model"], "gpt-5.4-mini")
        self.assertEqual(openai.return_value.responses.create.call_args.kwargs["reasoning"], {"effort": "low"})
        user_prompt = openai.return_value.responses.create.call_args.kwargs["input"][1]["content"]
        self.assertNotIn("jane@example.org", user_prompt)
        self.assertNotIn("sk_secretvalue12345", user_prompt)
        self.assertIn("[redacted-email]", user_prompt)

    def test_generate_checklist_prompt_uses_requested_language(self):
        output_text = json.dumps(
            {
                "event_title": "Lokakarya Aman",
                "source_note": "Dibuat.",
                "sections": [{"title": "Daftar cepat", "items": ["Bawa air."]}],
            }
        )

        with patch("resources.checklist_generator.OpenAI") as openai:
            openai.return_value.responses.create.return_value = SimpleNamespace(output_text=output_text)
            generate_checklist_payload("Lokakarya marimba untuk siswa desa.", "sk-test", language="id")

        system_prompt = openai.return_value.responses.create.call_args.kwargs["input"][0]["content"]
        self.assertIn("Indonesian/Bahasa Indonesia", system_prompt)
        self.assertNotIn("Use English unless", system_prompt)

    @override_settings(CHECKLIST_MAX_CONCEPT_NOTE_CHARS=20)
    def test_generate_checklist_bounds_prompt_size(self):
        output_text = json.dumps(
            {
                "event_title": "Safe Workshop",
                "source_note": "Generated.",
                "sections": [{"title": "Quick checklist", "items": ["Bring water."]}],
            }
        )

        with patch("resources.checklist_generator.OpenAI") as openai:
            openai.return_value.responses.create.return_value = SimpleNamespace(output_text=output_text)
            generate_checklist_payload("A" * 50, "sk-test")

        user_prompt = openai.return_value.responses.create.call_args.kwargs["input"][1]["content"]
        self.assertIn("A" * 20, user_prompt)
        self.assertNotIn("A" * 21, user_prompt)

    def test_generate_checklist_endpoint_requires_queue_endpoint(self):
        self.client.login(username="volunteer", password="test-password")
        self.user.user_permissions.add(Permission.objects.get(codename="can_generate_checklist"))

        response = self.client.post(reverse("api-checklist-generate"))

        self.assertEqual(response.status_code, 410)
        self.assertIn("queued", response.json()["detail"])

    def test_openai_api_key_is_encrypted_at_rest(self):
        settings = OpenAISettings.objects.create(api_key="sk-test")

        with connection.cursor() as cursor:
            cursor.execute("SELECT api_key FROM resources_openaisettings WHERE id = %s", [settings.id])
            stored_value = cursor.fetchone()[0]

        self.assertNotEqual(stored_value, "sk-test")
        self.assertTrue(stored_value.startswith("fernet$"))
        self.assertEqual(OpenAISettings.objects.get(id=settings.id).api_key, "sk-test")

    def test_generate_checklist_rejects_invalid_docx(self):
        self.client.login(username="volunteer", password="test-password")
        self.user.user_permissions.add(Permission.objects.get(codename="can_generate_checklist"))
        upload = SimpleUploadedFile(
            "concept-note.docx",
            b"not a docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        with self.settings(CELERY_TASK_ALWAYS_EAGER=True):
            response = self.client.post(reverse("api-checklist-job-create"), checklist_job_payload(upload))

        self.assertEqual(response.status_code, 202)
        job = ChecklistJob.objects.get()
        self.assertEqual(job.status, ChecklistJob.Status.ERROR)
        self.assertIn("valid DOCX", job.error_message)

    @override_settings(CHECKLIST_MAX_UPLOAD_BYTES=8)
    def test_generate_checklist_rejects_oversized_docx(self):
        self.client.login(username="volunteer", password="test-password")
        self.user.user_permissions.add(Permission.objects.get(codename="can_generate_checklist"))
        upload = SimpleUploadedFile(
            "large.docx",
            build_readable_docx("Large concept note."),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        with self.settings(CELERY_TASK_ALWAYS_EAGER=True):
            response = self.client.post(reverse("api-checklist-job-create"), checklist_job_payload(upload))

        self.assertEqual(response.status_code, 202)
        job = ChecklistJob.objects.get()
        self.assertEqual(job.status, ChecklistJob.Status.ERROR)
        self.assertIn("Event Concept Note must be", job.error_message)

    @override_settings(DOCX_MAX_COMPRESSION_RATIO=1)
    def test_generate_checklist_rejects_high_compression_docx(self):
        self.client.login(username="volunteer", password="test-password")
        self.user.user_permissions.add(Permission.objects.get(codename="can_generate_checklist"))
        upload = SimpleUploadedFile(
            "compressed.docx",
            build_readable_docx("A" * 2000),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        with self.settings(CELERY_TASK_ALWAYS_EAGER=True):
            response = self.client.post(reverse("api-checklist-job-create"), checklist_job_payload(upload))

        self.assertEqual(response.status_code, 202)
        job = ChecklistJob.objects.get()
        self.assertEqual(job.status, ChecklistJob.Status.ERROR)
        self.assertIn("compression ratio", job.error_message)

    def test_generate_checklist_rejects_empty_docx(self):
        self.client.login(username="volunteer", password="test-password")
        self.user.user_permissions.add(Permission.objects.get(codename="can_generate_checklist"))
        upload = SimpleUploadedFile(
            "empty.docx",
            build_readable_docx("", include_heading=False),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        with self.settings(CELERY_TASK_ALWAYS_EAGER=True):
            response = self.client.post(reverse("api-checklist-job-create"), checklist_job_payload(upload))

        self.assertEqual(response.status_code, 202)
        job = ChecklistJob.objects.get()
        self.assertEqual(job.status, ChecklistJob.Status.ERROR)
        self.assertIn("does not contain readable text", job.error_message)

    def test_generate_checklist_requires_openai_key(self):
        self.client.login(username="volunteer", password="test-password")
        self.user.user_permissions.add(Permission.objects.get(codename="can_generate_checklist"))
        upload = SimpleUploadedFile(
            "concept-note.docx",
            build_readable_docx("Marimba workshop for village students."),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        with self.settings(CELERY_TASK_ALWAYS_EAGER=True):
            response = self.client.post(reverse("api-checklist-job-create"), checklist_job_payload(upload))

        self.assertEqual(response.status_code, 202)
        job = ChecklistJob.objects.get()
        self.assertEqual(job.status, ChecklistJob.Status.ERROR)
        self.assertIn("OpenAI API key is not configured", job.error_message)

    def test_create_checklist_job_persists_pdf_for_preview_and_download(self):
        self.client.login(username="volunteer", password="test-password")
        self.user.user_permissions.add(Permission.objects.get(codename="can_generate_checklist"))
        OpenAISettings.objects.create(api_key="sk-test", checklist_queue_timeout_minutes=45)
        upload = SimpleUploadedFile(
            "concept-note.docx",
            build_readable_docx("Marimba workshop for village students."),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        output_text = json.dumps(
            {
                "event_title": "Marimba Workshop",
                "source_note": "Generated from uploaded Event Concept Note.",
                "sections": [{"title": "Quick checklist", "items": ["Bring water and notebook."]}],
            }
        )

        with self.settings(CELERY_TASK_ALWAYS_EAGER=True):
            with patch("resources.checklist_generator.OpenAI") as openai:
                openai.return_value.responses.create.return_value = SimpleNamespace(output_text=output_text)
                response = self.client.post(reverse("api-checklist-job-create"), checklist_job_payload(upload))

        self.assertEqual(response.status_code, 202)
        job_payload = response.json()["job"]
        uuid.UUID(job_payload["id"])
        self.assertEqual(job_payload["status"], "done")
        self.assertTrue(job_payload["previewUrl"])
        self.assertTrue(job_payload["downloadUrl"])

        list_response = self.client.get(reverse("api-checklist-job-list"))
        self.assertEqual(list_response.status_code, 200)
        self.assertEqual(len(list_response.json()["jobs"]), 1)

        preview_response = self.client.get(reverse("api-checklist-job-preview", args=[job_payload["id"]]))
        self.assertEqual(preview_response.status_code, 200)
        self.assertEqual(preview_response["Content-Type"], "application/pdf")
        self.assertIn("inline", preview_response["Content-Disposition"])

        download_response = self.client.get(reverse("api-checklist-job-download", args=[job_payload["id"]]))
        self.assertEqual(download_response.status_code, 200)
        self.assertIn("attachment", download_response["Content-Disposition"])

    def test_checklist_job_list_hides_expired_jobs(self):
        self.client.login(username="volunteer", password="test-password")
        self.user.user_permissions.add(Permission.objects.get(codename="can_generate_checklist"))
        expired_job = ChecklistJob.objects.create(
            user=self.user,
            input_filename="expired.docx",
            concept_note=SimpleUploadedFile("expired.docx", build_readable_docx(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            output_filename="expired.pdf",
            status=ChecklistJob.Status.DONE,
            expires_at=timezone.now() - timezone.timedelta(minutes=1),
        )
        expired_job.generated_pdf.save("expired.pdf", ContentFile(build_simple_pdf("Expired", "Expired")), save=True)

        response = self.client.get(reverse("api-checklist-job-list"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["jobs"], [])

    def test_cleanup_deletes_expired_job_files(self):
        expired_job = ChecklistJob.objects.create(
            user=self.user,
            input_filename="expired.docx",
            concept_note=SimpleUploadedFile("expired.docx", build_readable_docx(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            output_filename="expired.pdf",
            status=ChecklistJob.Status.DONE,
            expires_at=timezone.now() - timezone.timedelta(minutes=1),
        )
        expired_job.generated_pdf.save("expired.pdf", ContentFile(build_simple_pdf("Expired", "Expired")), save=True)
        concept_path = expired_job.concept_note.path
        pdf_path = expired_job.generated_pdf.path

        deleted = delete_expired_checklist_jobs()

        self.assertEqual(deleted, 1)
        self.assertFalse(ChecklistJob.objects.filter(id=expired_job.id).exists())
        self.assertFalse(Path(concept_path).exists())
        self.assertFalse(Path(pdf_path).exists())


@override_settings(CHAT_VECTOR_SEARCH_ENABLED=False)
class EducationChatApiTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="volunteer", password="test-password")
        self.staff = User.objects.create_user(username="staff", password="test-password", is_staff=True)
        OpenAISettings.objects.create(api_key="sk-test")
        self.educational = Resource.objects.create(
            title="Tsunami Education Guide",
            description="School evacuation routes and safe assembly points.",
            category=Resource.Category.EDUCATIONAL,
            pdf_file=SimpleUploadedFile(
                "tsunami-guide.docx",
                build_readable_docx("During shaking, drop, cover, and hold. Move to signed evacuation routes."),
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        )
        self.staff_only = Resource.objects.create(
            title="Staff Education Note",
            description="Internal contact tree.",
            category=Resource.Category.EDUCATIONAL,
            access_level=Resource.AccessLevel.STAFF,
            pdf_file=SimpleUploadedFile(
                "staff-note.docx",
                build_readable_docx("Staff only emergency contacts."),
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        )

    def test_anonymous_users_cannot_use_chat(self):
        response = self.client.post(reverse("api-chat-reply"), data=json.dumps({"message": "What should I read?"}), content_type="application/json")

        self.assertEqual(response.status_code, 401)

    def test_chat_uses_langgraph_chat_model_with_visible_educational_context(self):
        self.client.login(username="volunteer", password="test-password")

        with patch("resources.education_chat.ChatOpenAI") as chat_model:
            chat_model.return_value.invoke.return_value = SimpleNamespace(content="Use the Tsunami Education Guide.")
            response = self.client.post(
                reverse("api-chat-reply"),
                data=json.dumps({"message": "What should volunteers do during shaking?", "history": []}),
                content_type="application/json",
            )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["model"], "gpt-5.4-mini")
        self.assertEqual(payload["provider"], "OpenAI")
        self.assertEqual(payload["reasoningEffort"], "medium")
        self.assertEqual(payload["reply"], "Use the Tsunami Education Guide.")
        self.assertEqual(payload["sources"][0]["title"], "Tsunami Education Guide")
        chat_model.assert_called_once()
        self.assertEqual(chat_model.call_args.kwargs["model"], "gpt-5.4-mini")
        self.assertEqual(chat_model.call_args.kwargs["reasoning"], {"effort": "medium"})
        system_message = chat_model.return_value.invoke.call_args.args[0][0].content
        self.assertIn("Tsunami Education Guide", system_message)
        self.assertIn("Answer in English.", system_message)
        self.assertNotIn("Staff Education Note", system_message)

    def test_chat_rejects_unsupported_language(self):
        self.client.login(username="volunteer", password="test-password")

        response = self.client.post(
            reverse("api-chat-reply"),
            data=json.dumps({"message": "What should volunteers do during shaking?", "language": "fr"}),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["detail"], "Unsupported language.")

    def test_chat_prompt_uses_requested_language(self):
        self.client.login(username="volunteer", password="test-password")

        with patch("resources.education_chat.ChatOpenAI") as chat_model:
            chat_model.return_value.invoke.return_value = SimpleNamespace(content="Gunakan Panduan Edukasi Tsunami.")
            response = self.client.post(
                reverse("api-chat-reply"),
                data=json.dumps(
                    {
                        "message": "Apa yang harus dilakukan relawan saat guncangan?",
                        "history": [],
                        "language": "id",
                    }
                ),
                content_type="application/json",
            )

        self.assertEqual(response.status_code, 200)
        system_message = chat_model.return_value.invoke.call_args.args[0][0].content
        self.assertIn("Answer in Indonesian/Bahasa Indonesia.", system_message)

    @override_settings(CHAT_MAX_GROUNDING_SOURCES=10, CHAT_MAX_LINKED_SOURCES_PER_WEBSITE=2)
    def test_chat_searches_linked_pages_and_documents_from_website_source(self):
        ChatSource.objects.create(title="GMLS Education", url="https://example.org/education", created_by=self.staff)
        nav_links = "".join(f'<a href="/nav/{index}">Navigation {index}</a>' for index in range(12))
        responses = {
            "https://example.org/education": FakeWebsiteResponse(
                f"""
                <html>
                  <head><title>Education Home</title></head>
                  <body>
                    {nav_links}
                    <a href="/guides/tsunami">Tsunami linked guide</a>
                    <a href="/files/field-kit.docx">Field kit</a>
                  </body>
                </html>
                """
            ),
            "https://example.org/guides/tsunami": FakeWebsiteResponse(
                """
                <html>
                  <head><title>Tsunami Linked Guide</title></head>
                  <body>Linked guide says volunteers should move inland after official evacuation direction.</body>
                </html>
                """
            ),
            "https://example.org/files/field-kit.docx": FakeWebsiteResponse(
                build_readable_docx("Field kits include whistles, printed maps, and a charged phone."),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        }

        def open_website(request, timeout=None):
            return responses[request.full_url]

        self.client.login(username="volunteer", password="test-password")
        with patch("resources.education_chat.URL_OPENER.open", side_effect=open_website):
            with patch("resources.education_chat.ChatOpenAI") as chat_model:
                chat_model.return_value.invoke.return_value = SimpleNamespace(content="Use the linked website sources.")
                response = self.client.post(
                    reverse("api-chat-reply"),
                    data=json.dumps({"message": "What do linked sources say about evacuation and field kits?"}),
                    content_type="application/json",
                )

        self.assertEqual(response.status_code, 200)
        system_message = chat_model.return_value.invoke.call_args.args[0][0].content
        self.assertIn("Linked guide says volunteers should move inland", system_message)
        self.assertIn("Field kits include whistles", system_message)
        source_titles = {source["title"] for source in response.json()["sources"]}
        self.assertIn("GMLS Education: Tsunami Linked Guide", source_titles)
        self.assertIn("GMLS Education: field-kit.docx", source_titles)

    @override_settings(CHAT_MAX_GROUNDING_SOURCES=5, CHAT_SOURCE_SNIPPET_CHARS=700)
    def test_chat_uses_listed_webpage_content_after_navigation_for_generic_questions(self):
        ChatSource.objects.create(title="BMKG Education", url="https://example.org/education", created_by=self.staff)
        page_text = " ".join(f"Navigation item {index}" for index in range(120))
        responses = {
            "https://example.org/education": FakeWebsiteResponse(
                f"""
                <html>
                  <head><title>Edukasi Gempa Bumi dan Tsunami</title></head>
                  <body>
                    <nav>{page_text}</nav>
                    <main>
                      <h1>Buku</h1>
                      <article>
                        <h2>Buku Saku Gempabumi dan Tsunami</h2>
                        <p>Panduan edukatif mengenai penyebab, dampak, dan langkah mitigasi sebelum, saat,
                        dan setelah bencana terjadi.</p>
                      </article>
                      <article>
                        <h2>Petunjuk Keselamatan Gempabumi dan Tsunami Bagi Penyandang Disabilitas Netra</h2>
                        <p>Panduan evakuasi, titik kumpul ramah disabilitas, dan pertolongan pertama.</p>
                      </article>
                    </main>
                  </body>
                </html>
                """
            ),
        }

        def open_website(request, timeout=None):
            return responses[request.full_url]

        self.client.login(username="volunteer", password="test-password")
        with patch("resources.education_chat.URL_OPENER.open", side_effect=open_website):
            with patch("resources.education_chat.ChatOpenAI") as chat_model:
                chat_model.return_value.invoke.return_value = SimpleNamespace(content="The page lists tsunami education books.")
                response = self.client.post(
                    reverse("api-chat-reply"),
                    data=json.dumps({"message": "What content is listed on this webpage?"}),
                    content_type="application/json",
                )

        self.assertEqual(response.status_code, 200)
        system_message = chat_model.return_value.invoke.call_args.args[0][0].content
        self.assertIn("Buku Saku Gempabumi dan Tsunami", system_message)
        self.assertIn("Petunjuk Keselamatan Gempabumi", system_message)

    def test_chat_answers_basic_earthquake_question_from_portal_glossary(self):
        self.client.login(username="volunteer", password="test-password")

        with patch("resources.education_chat.ChatOpenAI") as chat_model:
            chat_model.return_value.invoke.return_value = SimpleNamespace(
                content="An earthquake is ground shaking caused by sudden energy release inside the Earth."
            )
            response = self.client.post(
                reverse("api-chat-reply"),
                data=json.dumps({"message": "What is earthquake?"}),
                content_type="application/json",
            )

        self.assertEqual(response.status_code, 200)
        system_message = chat_model.return_value.invoke.call_args.args[0][0].content
        self.assertIn("Volunteer education glossary: Earthquake", system_message)
        self.assertIn("Earthquake (gempa bumi/gempabumi) is ground shaking", system_message)
        source_titles = [source["title"] for source in response.json()["sources"]]
        self.assertIn("Volunteer education glossary: Earthquake", source_titles)

    def test_chat_corrects_misspelled_earthquake_question_from_portal_glossary(self):
        self.client.login(username="volunteer", password="test-password")

        with patch("resources.education_chat.ChatOpenAI") as chat_model:
            chat_model.return_value.invoke.return_value = SimpleNamespace(
                content="An earthquake is ground shaking caused by sudden energy release inside the Earth."
            )
            response = self.client.post(
                reverse("api-chat-reply"),
                data=json.dumps({"message": "waht is earthqauke"}),
                content_type="application/json",
            )

        self.assertEqual(response.status_code, 200)
        system_message = chat_model.return_value.invoke.call_args.args[0][0].content
        self.assertIn("Volunteer education glossary: Earthquake", system_message)
        self.assertIn("Earthquake (gempa bumi/gempabumi) is ground shaking", system_message)
        self.assertIn("gempa bumi", education_chat._query_terms("waht is earthqauke"))
        source_titles = [source["title"] for source in response.json()["sources"]]
        self.assertIn("Volunteer education glossary: Earthquake", source_titles)

    @override_settings(CHAT_MAX_GROUNDING_SOURCES=8, CHAT_MAX_LINKED_SOURCES_PER_WEBSITE=1)
    def test_chat_expands_english_earthquake_query_to_indonesian_linked_sources(self):
        ChatSource.objects.create(title="BMKG Education", url="https://example.org/education", created_by=self.staff)
        nav_links = "".join(f'<a href="/nav/{index}">Navigation {index}</a>' for index in range(12))
        responses = {
            "https://example.org/education": FakeWebsiteResponse(
                f"""
                <html>
                  <head><title>Edukasi Bencana</title></head>
                  <body>
                    {nav_links}
                    <a href="/mitigasi/antisipasi-gempabumi">Antisipasi Gempabumi</a>
                  </body>
                </html>
                """
            ),
            "https://example.org/mitigasi/antisipasi-gempabumi": FakeWebsiteResponse(
                """
                <html>
                  <head><title>Antisipasi Gempabumi</title></head>
                  <body>Gempa bumi adalah getaran tanah akibat pelepasan energi secara tiba-tiba.</body>
                </html>
                """
            ),
        }

        def open_website(request, timeout=None):
            return responses[request.full_url]

        self.client.login(username="volunteer", password="test-password")
        with patch("resources.education_chat.URL_OPENER.open", side_effect=open_website):
            with patch("resources.education_chat.ChatOpenAI") as chat_model:
                chat_model.return_value.invoke.return_value = SimpleNamespace(content="Gempa bumi means earthquake.")
                response = self.client.post(
                    reverse("api-chat-reply"),
                    data=json.dumps({"message": "What is earthquake?"}),
                    content_type="application/json",
                )

        self.assertEqual(response.status_code, 200)
        system_message = chat_model.return_value.invoke.call_args.args[0][0].content
        self.assertIn("Gempa bumi adalah getaran tanah", system_message)
        source_titles = {source["title"] for source in response.json()["sources"]}
        self.assertIn("BMKG Education: Antisipasi Gempabumi", source_titles)

    def test_chat_uses_gpt_even_when_no_sources_are_available(self):
        Resource.objects.all().delete()
        self.client.login(username="volunteer", password="test-password")

        with patch("resources.education_chat.ChatOpenAI") as chat_model:
            chat_model.return_value.invoke.return_value = SimpleNamespace(
                content="No approved source material was found for that question."
            )
            response = self.client.post(
                reverse("api-chat-reply"),
                data=json.dumps({"message": "What is the training schedule?"}),
                content_type="application/json",
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["reply"], "No approved source material was found for that question.")
        self.assertEqual(response.json()["sources"], [])
        chat_model.assert_called_once()
        system_message = chat_model.return_value.invoke.call_args.args[0][0].content
        self.assertIn("No approved portal educational files", system_message)

    def test_chat_provider_errors_are_returned_as_setup_errors(self):
        self.client.login(username="volunteer", password="test-password")

        with patch("resources.education_chat.ChatOpenAI") as chat_model:
            chat_model.return_value.invoke.side_effect = RuntimeError("model unavailable")
            response = self.client.post(
                reverse("api-chat-reply"),
                data=json.dumps({"message": "What should volunteers do during shaking?"}),
                content_type="application/json",
            )

        self.assertEqual(response.status_code, 503)
        self.assertIn("GPT request failed for gpt-5.4-mini", response.json()["detail"])

    def test_chat_requires_openai_key_when_sources_exist(self):
        OpenAISettings.objects.update(api_key="")
        self.client.login(username="volunteer", password="test-password")

        response = self.client.post(
            reverse("api-chat-reply"),
            data=json.dumps({"message": "What should volunteers do during shaking?"}),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 503)
        self.assertIn("OpenAI API key", response.json()["detail"])

    def test_staff_can_manage_chat_website_sources(self):
        self.client.login(username="staff", password="test-password")

        create_response = self.client.post(
            reverse("api-chat-source-create"),
            data=json.dumps({"title": "GMLS Education", "url": "https://example.org/education"}),
            content_type="application/json",
        )

        self.assertEqual(create_response.status_code, 201)
        source_id = create_response.json()["source"]["id"]
        self.assertTrue(ChatSource.objects.filter(title="GMLS Education").exists())

        list_response = self.client.get(reverse("api-chat-source-list"))
        self.assertEqual(list_response.status_code, 200)
        self.assertEqual(len(list_response.json()["sources"]), 1)

        update_response = self.client.post(
            reverse("api-chat-source-update", args=[source_id]),
            data=json.dumps({"isActive": False}),
            content_type="application/json",
        )
        self.assertEqual(update_response.status_code, 200)
        self.assertFalse(ChatSource.objects.get(public_id=source_id).is_active)

        delete_response = self.client.post(reverse("api-chat-source-delete", args=[source_id]))
        self.assertEqual(delete_response.status_code, 200)
        self.assertFalse(ChatSource.objects.exists())

    def test_non_staff_cannot_manage_chat_website_sources(self):
        self.client.login(username="volunteer", password="test-password")

        response = self.client.get(reverse("api-chat-source-list"))

        self.assertEqual(response.status_code, 403)

    def test_chat_website_source_rejects_private_urls(self):
        self.client.login(username="staff", password="test-password")

        response = self.client.post(
            reverse("api-chat-source-create"),
            data=json.dumps({"url": "http://127.0.0.1/internal"}),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("public", response.json()["detail"])


class EducationChatHybridRetrievalTests(TestCase):
    @override_settings(CHAT_VECTOR_SEARCH_ENABLED=True, CHAT_VECTOR_MAX_CANDIDATES=10, OPENAI_EMBEDDING_MODEL="text-embedding-3-small")
    def test_vector_source_ranks_use_langchain_vector_store(self):
        sources = [
            education_chat.GroundingSource(
                title="String earthquake source",
                kind="Approved website",
                locator="https://example.org/string",
                snippet="Earthquake preparedness checklist.",
                score=20,
            ),
            education_chat.GroundingSource(
                title="Semantic gempabumi source",
                kind="Linked website source",
                locator="https://example.org/gempabumi",
                snippet="Gempa bumi adalah getaran tanah akibat pelepasan energi secara tiba-tiba.",
                score=3,
            ),
        ]
        captured = {}

        class FakeVectorStore:
            def __init__(self, documents):
                self.documents = documents

            @classmethod
            def from_documents(cls, documents, embedding):
                captured["documents"] = documents
                captured["embedding"] = embedding
                return cls(documents)

            def similarity_search_with_score(self, question, k):
                captured["question"] = question
                captured["k"] = k
                return [(self.documents[1], 0.95), (self.documents[0], 0.25)]

        with patch("resources.education_chat.OpenAIEmbeddings") as embeddings:
            embeddings.return_value = object()
            with patch("resources.education_chat.InMemoryVectorStore", FakeVectorStore):
                ranks = education_chat._vector_source_ranks(sources, "What is earthquake?", "sk-real")

        embeddings.assert_called_once_with(model="text-embedding-3-small", api_key="sk-real", timeout=60)
        self.assertEqual(captured["question"], "What is earthquake?")
        self.assertEqual(captured["k"], 2)
        self.assertIn("Semantic gempabumi source", captured["documents"][1].page_content)
        self.assertEqual(ranks, {1: 1, 0: 2})

    @override_settings(CHAT_HYBRID_STRING_WEIGHT=1, CHAT_HYBRID_VECTOR_WEIGHT=100, CHAT_HYBRID_RRF_K=1)
    def test_hybrid_rank_sources_merges_vector_and_string_rankings(self):
        sources = [
            education_chat.GroundingSource(
                title="Exact string match",
                kind="Approved website",
                locator="https://example.org/string",
                snippet="The word earthquake appears here, but this chunk is only a navigation item.",
                score=100,
            ),
            education_chat.GroundingSource(
                title="Semantic bilingual match",
                kind="Linked website source",
                locator="https://example.org/gempabumi",
                snippet="Gempa bumi adalah getaran tanah akibat pelepasan energi secara tiba-tiba.",
                score=1,
            ),
        ]

        with patch("resources.education_chat._vector_source_ranks", return_value={1: 1, 0: 2}):
            ranked = education_chat._hybrid_rank_sources(sources, "What is earthquake?", "sk-real")

        self.assertEqual(ranked[0].title, "Semantic bilingual match")
        self.assertGreater(ranked[0].score, ranked[1].score)


class ThrottleTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="volunteer", password="test-password")

    @override_settings(LOGIN_THROTTLE_LIMIT=1)
    def test_api_login_is_throttled(self):
        first = self.client.post(
            reverse("api-login"),
            data=json.dumps({"username": "volunteer", "password": "wrong"}),
            content_type="application/json",
        )
        second = self.client.post(
            reverse("api-login"),
            data=json.dumps({"username": "volunteer", "password": "wrong"}),
            content_type="application/json",
        )

        self.assertEqual(first.status_code, 400)
        self.assertEqual(second.status_code, 429)
        self.assertTrue(ThrottleRecord.objects.filter(scope="api_login").exists())

    @override_settings(ADMIN_LOGIN_THROTTLE_LIMIT=1)
    def test_admin_login_is_throttled(self):
        first = self.client.post("/admin/login/", {"username": "admin", "password": "wrong"})
        second = self.client.post("/admin/login/", {"username": "admin", "password": "wrong"})

        self.assertNotEqual(first.status_code, 429)
        self.assertEqual(second.status_code, 429)

    @override_settings(TRUSTED_PROXY_IPS=["10.0.0.10"], CLIENT_IP_HEADER="HTTP_X_FORWARDED_FOR")
    def test_client_ip_uses_forwarded_header_only_from_trusted_proxy(self):
        trusted_request = SimpleNamespace(
            META={"REMOTE_ADDR": "10.0.0.10", "HTTP_X_FORWARDED_FOR": "203.0.113.5, 10.0.0.10"}
        )
        untrusted_request = SimpleNamespace(
            META={"REMOTE_ADDR": "198.51.100.10", "HTTP_X_FORWARDED_FOR": "203.0.113.5"}
        )

        self.assertEqual(client_ip(trusted_request), "203.0.113.5")
        self.assertEqual(client_ip(untrusted_request), "198.51.100.10")
